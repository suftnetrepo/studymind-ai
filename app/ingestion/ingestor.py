"""
Document ingestion pipeline — Sprint 3 extended.
DOC-04: Every chunk tagged with module_id, course_code, semester_id,
        week_number, visibility, owner_id, is_latest, doc_version,
        is_current_semester, institution_id, lecturer_id.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import Document as LIDocument

from app.agents.llm_factory import get_embed_model
from app.config import get_settings
from app.logging_config import get_logger
from app.retrieval.typesense_client import get_typesense_client, upsert_chunks

log = get_logger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".csv", ".docx"}


def _extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return content.decode("utf-8", errors="replace")
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            return "\n\n".join(p.extract_text() or "" for p in reader.pages)
        except ImportError:
            return content.decode("utf-8", errors="replace")
    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text)
        except ImportError:
            return content.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {ext}")


def _chunk_text(text: str, filename: str, document_id: str):
    settings = get_settings()
    splitter = SentenceSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    base_doc = LIDocument(
        text=text,
        metadata={"filename": filename, "document_id": document_id},
    )
    return splitter.get_nodes_from_documents([base_doc])


class DocumentIngestor:
    """
    Stateless ingestor.
    Returns pure data — no ORM objects.
    Caller persists to PostgreSQL.
    """

    def __init__(self) -> None:
        self._embed_model = get_embed_model()
        self._ts_client   = get_typesense_client()

    def ingest(
        self,
        filename: str,
        content: bytes,
        document_id: str,
        # DOC-04: Scoping metadata
        module_id: str = "",
        course_code: str = "",
        semester_id: str = "",
        week_id: str = "",
        week_number: int = 0,
        visibility: str = "personal",
        owner_id: str = "",
        doc_version: int = 1,
        is_current_semester: bool = True,
        institution_id: str = "",
        lecturer_id: str = "",
    ) -> dict:
        """
        Full ingestion flow.

        Returns:
            {
                status:      "indexed" | "failed",
                chunk_count: int,
                error:       str | None,
                chunks:      list[dict],
                indexed_at:  datetime | None,
            }
        """
        log.info("ingestion_start", filename=filename, doc_id=document_id)

        try:
            text = _extract_text(filename, content)
            if not text.strip():
                raise ValueError("Extracted text is empty — file may be image-only or corrupt")

            nodes = _chunk_text(text, filename, document_id)
            log.info("chunks_created", count=len(nodes), doc_id=document_id)

            texts_to_embed = [node.get_content() for node in nodes]
            embeddings     = self._embed_model.get_text_embedding_batch(
                texts_to_embed, show_progress=False
            )

            ts_docs:   list[dict] = []
            db_chunks: list[dict] = []

            for idx, (node, embedding) in enumerate(zip(nodes, embeddings)):
                chunk_text  = node.get_content()
                ts_id       = f"{document_id}__{idx}"
                token_count = len(chunk_text.split())

                # DOC-04: Full scoping metadata on every chunk
                ts_docs.append({
                    "id":                   ts_id,
                    "content":              chunk_text,
                    "embedding":            embedding,
                    "chunk_index":          idx,
                    "token_count":          token_count,
                    # Document identity
                    "document_id":          document_id,
                    "filename":             filename,
                    # Module / course
                    "module_id":            module_id or "none",
                    "course_code":          course_code.upper() if course_code else "NONE",
                    # Semester
                    "semester_id":          semester_id or "none",
                    "is_current_semester":  is_current_semester,
                    # Week
                    "week_id":              week_id or "none",
                    "week_number":          week_number,
                    # Access control
                    "visibility":           visibility,
                    "owner_id":             owner_id or "none",
                    # Version
                    "is_latest":            True,
                    "doc_version":          doc_version,
                    # Institution
                    "institution_id":       institution_id or "none",
                    "lecturer_id":          lecturer_id or "none",
                })

                db_chunks.append({
                    "typesense_id":   ts_id,
                    "chunk_index":    idx,
                    "content":        chunk_text,
                    "token_count":    token_count,
                    "chunk_metadata": {
                        **node.metadata,
                        "module_id":   module_id,
                        "course_code": course_code,
                        "week_number": week_number,
                        "visibility":  visibility,
                    },
                })

            upsert_chunks(self._ts_client, ts_docs)

            log.info("ingestion_complete", doc_id=document_id, chunks=len(db_chunks))
            return {
                "status":      "indexed",
                "chunk_count": len(db_chunks),
                "error":       None,
                "chunks":      db_chunks,
                "indexed_at":  datetime.now(timezone.utc),
            }

        except Exception as exc:
            log.error("ingestion_failed", doc_id=document_id, error=str(exc))
            return {
                "status":      "failed",
                "chunk_count": 0,
                "error":       str(exc),
                "chunks":      [],
                "indexed_at":  None,
            }


_ingestor: DocumentIngestor | None = None


def get_ingestor() -> DocumentIngestor:
    global _ingestor
    if _ingestor is None:
        _ingestor = DocumentIngestor()
    return _ingestor
